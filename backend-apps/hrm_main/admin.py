import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.contrib import admin
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import path

from hrm_main.models import (
    AttendanceStatusMaster,
    EmployeeMonthlySalary,
    EmployeeMonthlyAllowanceDetails,
)


@admin.register(AttendanceStatusMaster)
class AttendanceStatusMasterAdmin(admin.ModelAdmin):
    list_display = ('code', 'status_name', 'description')
    search_fields = ('code', 'status_name')
    list_filter = ('status_name',)
    ordering = ('code',)
    fields = ('code', 'status_name', 'description')


class EmployeeMonthlyAllowanceDetailsInline(admin.TabularInline):
    model = EmployeeMonthlyAllowanceDetails
    extra = 0
    readonly_fields = ('salary_component', 'components', 'monthly', 'yearly')
    fields = ('salary_component', 'components', 'monthly', 'yearly')
    can_delete = False

    def has_add_permission(self, request, obj=None):
        return False


@admin.register(EmployeeMonthlySalary)
class EmployeeMonthlySalaryAdmin(admin.ModelAdmin):
    change_list_template = 'admin/hrm/employeemonthlysalary_changelist.html'
    list_display = (
        'get_employee_code', 'get_employee_name', 'b_id',
        'net_pay_monthly', 'net_pay_yearly', 'ctc', 'created',
    )
    search_fields = (
        'employee__employee_code', 'employee__first_name',
        'employee__last_name', 'b_id',
    )
    list_filter = ('b_id',)
    readonly_fields = ('employee', 'net_pay_monthly', 'net_pay_yearly', 'ctc', 'b_id', 'created', 'modified')
    inlines = [EmployeeMonthlyAllowanceDetailsInline]

    def get_employee_code(self, obj):
        return obj.employee.employee_code if obj.employee else '-'
    get_employee_code.short_description = 'Employee Code'
    get_employee_code.admin_order_field = 'employee__employee_code'

    def get_employee_name(self, obj):
        if not obj.employee:
            return '-'
        return f"{obj.employee.first_name or ''} {obj.employee.last_name or ''}".strip()
    get_employee_name.short_description = 'Employee Name'
    get_employee_name.admin_order_field = 'employee__first_name'

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'import-salary-components/',
                self.admin_site.admin_view(self.import_salary_components_view),
                name='hrm_employeemonthlysalary_import_salary_components',
            ),
            path(
                'import-salary-components/sample/',
                self.admin_site.admin_view(self.download_sample_excel),
                name='hrm_employeemonthlysalary_import_sample',
            ),
        ]
        return custom_urls + urls

    def download_sample_excel(self, request):
        from hrm_master.models import SalaryComponents
        b_id = request.GET.get('b_id', '').strip()

        components = []
        if b_id:
            components = list(
                SalaryComponents.objects.filter(b_id=b_id).order_by('order').values_list('component', flat=True)
            )
        if not components:
            components = ['BASIC', 'TRAVEL', 'EDUCATION', 'HOUSING', 'TRANSPORT', 'OTHERS']

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Salary Components'

        header_fill = PatternFill(start_color='2E5F8A', end_color='2E5F8A', fill_type='solid')
        mandatory_fill = PatternFill(start_color='B80129', end_color='B80129', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=10)
        center = Alignment(horizontal='center', vertical='center')
        thin = Border(
            left=Side(style='thin', color='FFFFFF'),
            right=Side(style='thin', color='FFFFFF'),
            top=Side(style='thin', color='FFFFFF'),
            bottom=Side(style='thin', color='FFFFFF'),
        )

        headers = ['PAYROLL BRANCH', 'ECODE', 'EMPLOYEE NAME'] + components
        mandatory = {'PAYROLL BRANCH', 'ECODE', 'EMPLOYEE NAME'}

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.alignment = center
            cell.border = thin
            cell.fill = mandatory_fill if header in mandatory else header_fill
            ws.column_dimensions[cell.column_letter].width = max(18, len(header) + 4)

        # Two sample rows
        sample_rows = [
            ['1', 'E10001', 'Ahmed Raza Khan'] + ['0.00'] * len(components),
            ['1', 'E10002', 'Shoaib Ahmed'] + ['0.00'] * len(components),
        ]
        for row_data in sample_rows:
            ws.append(row_data)

        ws.row_dimensions[1].height = 22

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="salary_components_sample.xlsx"'
        return response

    def import_salary_components_view(self, request):
        import pandas as pd
        from hrm_master.models import EmployeeMaster, SalaryComponents

        context = {
            **self.admin_site.each_context(request),
            'title': 'Employee Salary Components Import',
            'opts': self.model._meta,
            'b_id': '',
            'summary': None,
            'results': [],
            'component_cols': [],
            'error': None,
        }

        if request.method == 'POST':
            excel_file = request.FILES.get('excel_file')

            if not excel_file:
                context['error'] = 'Please select an Excel file.'
                return render(request, 'admin/hrm/salary_components_import.html', context)

            try:
                df = pd.read_excel(excel_file, dtype=str)
                df = df.fillna('')
            except Exception as e:
                context['error'] = f'Failed to read Excel file: {e}'
                return render(request, 'admin/hrm/salary_components_import.html', context)

            if df.shape[1] < 4:
                context['error'] = 'Excel must have at least 4 columns: Branch, Employee Code, Employee Name, and at least one salary component.'
                return render(request, 'admin/hrm/salary_components_import.html', context)

            # Columns 0=b_id, 1=employee_code, 2=employee_name, 3+=components
            component_columns = list(df.columns[3:])

            # Build component map from all unique b_ids present in the file
            all_b_ids = df.iloc[:, 0].str.strip().unique().tolist()
            salary_components_map = {}
            for col in component_columns:
                for bid in all_b_ids:
                    sc = SalaryComponents.objects.filter(
                        component__iexact=col.strip(), b_id=bid
                    ).first()
                    if sc:
                        salary_components_map[col] = sc
                        break

            matched_cols = list(salary_components_map.keys())
            context['component_cols'] = matched_cols

            results = []
            new_count = 0
            updated_count = 0
            skipped_count = 0

            user_str = f"{request.user.first_name or ''} {request.user.last_name or ''}".strip() or str(request.user)

            try:
                with transaction.atomic():
                    for _, row in df.iterrows():
                        # Read b_id from column 0 of each row
                        b_id = str(row.iloc[0]).strip()
                        employee_code = str(row.iloc[1]).strip()
                        if not employee_code or employee_code.lower() in ('nan', 'none', ''):
                            skipped_count += 1
                            continue

                        employee = EmployeeMaster.objects.filter(
                            employee_code=employee_code, b_id=b_id
                        ).first()

                        if not employee:
                            # Try finding by employee_code alone, then fix the b_id
                            employee = EmployeeMaster.objects.filter(
                                employee_code=employee_code
                            ).first()
                            if employee:
                                old_b_id = employee.b_id
                                employee.b_id = b_id
                                employee.save(update_fields=['b_id'])
                            else:
                                results.append({
                                    'employee_code': employee_code,
                                    'employee_name': str(row.iloc[2]).strip(),
                                    'component_values': ['-'] * len(matched_cols),
                                    'net_pay_monthly': '-',
                                    'status': 'skipped',
                                    'reason': 'Employee not found',
                                })
                                skipped_count += 1
                                continue

                        existing = EmployeeMonthlySalary.objects.filter(
                            employee=employee, b_id=b_id
                        ).first()

                        if existing:
                            salary_record = existing
                            is_new = False
                        else:
                            salary_record = EmployeeMonthlySalary.objects.create(
                                employee=employee,
                                b_id=b_id,
                                user_created=user_str,
                                user_modified=user_str,
                            )
                            employee.is_sal_created = True
                            employee.save(update_fields=['is_sal_created'])
                            is_new = True

                        total_monthly = 0.0
                        component_values = []

                        for col in matched_cols:
                            sc = salary_components_map[col]
                            raw_val = str(row.get(col, '') or '').strip().replace(',', '')
                            try:
                                monthly_val = float(raw_val) if raw_val and raw_val.lower() not in ('nan', '-', '--', 'none', 'null', '') else 0.0
                            except ValueError:
                                monthly_val = 0.0

                            detail = EmployeeMonthlyAllowanceDetails.objects.filter(
                                gross_earnings=salary_record, salary_component=sc
                            ).first()

                            if detail:
                                detail.monthly = monthly_val
                                detail.yearly = round(monthly_val * 12, 2)
                                detail.components = sc.component
                                detail.user_modified = user_str
                                detail.save(update_fields=['monthly', 'yearly', 'components', 'user_modified'])
                            else:
                                EmployeeMonthlyAllowanceDetails.objects.create(
                                    gross_earnings=salary_record,
                                    salary_component=sc,
                                    components=sc.component,
                                    monthly=monthly_val,
                                    yearly=round(monthly_val * 12, 2),
                                    b_id=b_id,
                                    user_created=user_str,
                                    user_modified=user_str,
                                )

                            total_monthly += monthly_val
                            component_values.append(f"{monthly_val:,.2f}" if monthly_val else '-')

                        salary_record.net_pay_monthly = round(total_monthly, 2)
                        salary_record.net_pay_yearly = round(total_monthly * 12, 2)
                        salary_record.user_modified = user_str
                        salary_record.save(update_fields=['net_pay_monthly', 'net_pay_yearly', 'user_modified'])

                        if is_new:
                            new_count += 1
                        else:
                            updated_count += 1

                        results.append({
                            'employee_code': employee_code,
                            'employee_name': f"{employee.first_name or ''} {employee.last_name or ''}".strip(),
                            'component_values': component_values,
                            'net_pay_monthly': f"{salary_record.net_pay_monthly:,.2f}",
                            'status': 'created' if is_new else 'updated',
                            'reason': '',
                        })

            except Exception as e:
                context['error'] = f'Import failed: {e}'
                return render(request, 'admin/hrm/salary_components_import.html', context)

            context['summary'] = {
                'total': new_count + updated_count + skipped_count,
                'created': new_count,
                'updated': updated_count,
                'skipped': skipped_count,
            }
            context['results'] = results

        return render(request, 'admin/hrm/salary_components_import.html', context)
